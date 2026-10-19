import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import collector


class ExcelExportTests(unittest.TestCase):
    def test_automatic_readable_export_preserves_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "corpus.csv"
            fields = ["number", "username", "corpus", "post_text", "comment_page_url", "note"]
            text = "完整语境。" * 40
            with source.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=fields)
                writer.writeheader()
                for number in (1, 2):
                    writer.writerow(dict(number=number, username=f"user{number}", corpus="=文本不是公式", post_text=text, comment_page_url="https://weibo.cn/comment/test", note="虚构测试"))
            original = source.read_bytes()
            collector.write_frequency_outputs(str(source))
            self.assertEqual(source.read_bytes(), original)
            workbook = load_workbook(source.with_suffix(".xlsx"))
            try:
                self.assertEqual(workbook.sheetnames, ["语料表", "频率表"])
                sheet = workbook["语料表"]
                self.assertEqual([cell.value for cell in sheet[1]], ["编号", "关键词", "用户名", "全文语境", "来源", "备注"])
                self.assertEqual(sheet["D2"].value, text)
                self.assertTrue(sheet["D2"].alignment.wrap_text)
                self.assertEqual(sheet["B2"].data_type, "s")
                self.assertEqual(sheet["E2"].hyperlink.target, "https://weibo.cn/comment/test")
                self.assertEqual(sheet.freeze_panes, "D2")
                self.assertEqual(sheet.tables["CorpusRecords"].autoFilter.ref, "A1:F3")
                self.assertGreater(sheet.row_dimensions[2].height, 48)
                self.assertEqual(workbook["频率表"]["C2"].value, 2)
                self.assertEqual(workbook["频率表"]["D2"].value, 2)
            finally:
                workbook.close()
            # Re-export replaces the previous workbook instead of accumulating rows.
            collector.write_frequency_outputs(str(source))
            self.assertTrue(source.with_suffix(".xlsx").is_file())


if __name__ == "__main__":
    unittest.main()
