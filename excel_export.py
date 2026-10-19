"""Export collected CSV records as a readable, portable Excel workbook."""

import csv
import math
import os
import tempfile
import unicodedata
from pathlib import Path


def export_excel(csv_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    source = Path(csv_path)
    with source.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        fields = reader.fieldnames or []
        records = list(reader)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "语料表"
    columns = [
        ("编号", "number", 8),
        ("关键词", "corpus", 24),
        ("用户名", "username", 24),
        ("全文语境", "post_text", 90),
        ("来源", "comment_page_url", 16),
        ("备注", "note", 28),
    ]
    if "published_at" in fields or "published_time_raw" in fields:
        columns.insert(4, ("发布时间", "published_at", 28))
    sheet.append([label for label, _, _ in columns])

    def literal(cell, value):
        # Scraped text must remain text, including expressions beginning with '='.
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"

    for number, record in enumerate(records, start=2):
        height = 48
        for index, (_, key, width) in enumerate(columns, start=1):
            value = record.get(key, "")
            if key == "post_text":
                value = value or record.get("comment_text", "")
            elif key == "published_at":
                value = value or record.get("published_time_raw", "")
            elif key == "number":
                value = int(value) if str(value).isdigit() else number - 1
            cell = sheet.cell(number, index)
            literal(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if key == "comment_page_url":
                url = value or record.get("source", "")
                if url.startswith(("https://", "http://")):
                    cell.hyperlink = url
                    literal(cell, "查看原文")
                    value = "查看原文"
                    cell.font = Font(name="Arial", size=11, color="245CA6", underline="single")
            units = sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(value))
            lines = max(str(value).count("\n") + 1, math.ceil(units / max(1, width - 3)))
            height = max(height, 18 * lines + 12)
        sheet.row_dimensions[number].height = min(409, height)

    def style_table(target, widths, name):
        target.sheet_view.showGridLines = False
        target.freeze_panes = "D2" if target == sheet else "A2"
        target.sheet_view.zoomScale = 90
        for row in target:
            for cell in row:
                if not cell.hyperlink:
                    cell.font = Font(name="Arial", size=11, color="20242A")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="top")
        if target != sheet:
            for number in range(2, target.max_row + 1):
                target.row_dimensions[number].height = 24
        for cell in target[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="344B62")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target.row_dimensions[1].height = 30
        for index, width in enumerate(widths, start=1):
            target.column_dimensions[get_column_letter(index)].width = width
        if target.max_row > 1:
            table = Table(displayName=name, ref=target.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            target.add_table(table)
        target.print_options.horizontalCentered = True
        target.print_title_rows = "1:1"
        target.sheet_properties.pageSetUpPr.fitToPage = True
        target.page_setup.orientation = "landscape"
        target.page_setup.paperSize = target.PAPERSIZE_A3
        target.page_setup.fitToWidth = 1
        target.page_setup.fitToHeight = 0

    style_table(sheet, [width for _, _, width in columns], "CorpusRecords")
    frequency = workbook.create_sheet("频率表")
    frequency.append(["排名", "关键词", "频次", "独立用户名数"])
    counts = {}
    users = {}
    for record in records:
        corpus = record.get("corpus", "").strip()
        if not corpus:
            continue
        counts[corpus] = counts.get(corpus, 0) + 1
        users.setdefault(corpus, set())
        if record.get("username", "").strip():
            users[corpus].add(record["username"].strip())
    for rank, corpus in enumerate(sorted(counts, key=lambda item: (-counts[item], item)), start=1):
        frequency.append([rank, None, counts[corpus], len(users[corpus])])
        literal(frequency.cell(rank + 1, 2), corpus)
    style_table(frequency, [8, 32, 12, 20], "CorpusFrequency")

    output = source.with_suffix(".xlsx")
    descriptor, temporary = tempfile.mkstemp(dir=source.parent, suffix=".xlsx")
    os.close(descriptor)
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
        workbook.close()
    return output
